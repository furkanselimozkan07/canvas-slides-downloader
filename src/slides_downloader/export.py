"""Writing captured slides to PNG files, a PDF and a PPTX."""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

FORMATS = ("pdf", "pptx", "png")
PDF_RESOLUTION = 150
PPTX_WIDTH = Emu(12192000)  # 13.333 in, PowerPoint's 16:9 default
_BLANK_LAYOUT = 6


def save_pngs(images: Sequence[bytes], out_dir: Path) -> list[Path]:
    for stale in out_dir.glob("slide_*.png"):
        stale.unlink()
    paths = []
    for number, png in enumerate(images, 1):
        path = out_dir / f"slide_{number:03d}.png"
        path.write_bytes(png)
        paths.append(path)
    return paths


def save_pdf(images: Sequence[bytes], path: Path) -> Path:
    pages = [Image.open(io.BytesIO(png)).convert("RGB") for png in images]
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=PDF_RESOLUTION)
    return path


def save_pptx(images: Sequence[bytes], path: Path) -> Path:
    """One slide per image, each image filling the whole slide."""
    width, height = Image.open(io.BytesIO(images[0])).size
    prs = Presentation()
    prs.slide_width = PPTX_WIDTH
    prs.slide_height = Emu(int(PPTX_WIDTH * height / width))
    layout = prs.slide_layouts[_BLANK_LAYOUT]
    for png in images:
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(path)
    return path


def export(images: Sequence[bytes], out_dir: Path, name: str, formats: Iterable[str]) -> list[Path]:
    """Write ``images`` to ``out_dir`` in the requested formats; return the written paths."""
    if not images:
        raise ValueError("no images to export")
    formats = set(formats)
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    if "pdf" in formats:
        written.append(save_pdf(images, out_dir / f"{name}.pdf"))
    if "pptx" in formats:
        written.append(save_pptx(images, out_dir / f"{name}.pptx"))
    if "png" in formats:
        written.extend(save_pngs(images, out_dir))
    return written
